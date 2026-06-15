import io

import pytest
from openpyxl import Workbook as XlWorkbook, load_workbook


async def _make_sheet(client) -> str:
    wb = await client.post("/api/v1/workbooks", json={"name": "WB"})
    sheet = await client.post(
        f"/api/v1/workbooks/{wb.json()['id']}/sheets",
        json={"name": "S"},
    )
    return sheet.json()["id"]


def _bytes_workbook(rows: list[list]) -> bytes:
    wb = XlWorkbook()
    ws = wb.active
    ws.title = "Imported"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_xlsx_import_creates_sheets_and_cells(client):
    wb_resp = await client.post("/api/v1/workbooks", json={"name": "WB"})
    wb_id = wb_resp.json()["id"]
    xlsx = _bytes_workbook([["name", "qty"], ["apple", 3], ["banana", 5]])
    response = await client.post(
        f"/api/v1/workbooks/{wb_id}/import/xlsx",
        files={"file": ("data.xlsx", io.BytesIO(xlsx),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    sheets = response.json()
    assert len(sheets) == 1
    sheet_id = sheets[0]["id"]
    cells = (await client.get(f"/api/v1/sheets/{sheet_id}/cells")).json()
    by = {(c["row"], c["column"]): c for c in cells}
    assert by[(0, 0)]["value"] == "name"
    assert by[(1, 1)]["value"] == "3"
    assert by[(1, 1)]["data_type"] == "number"


@pytest.mark.asyncio
async def test_xlsx_import_preserves_formulas(client):
    wb_resp = await client.post("/api/v1/workbooks", json={"name": "WB"})
    wb_id = wb_resp.json()["id"]
    wb = XlWorkbook()
    ws = wb.active
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=SUM(A1:A2)"
    buf = io.BytesIO()
    wb.save(buf)
    response = await client.post(
        f"/api/v1/workbooks/{wb_id}/import/xlsx",
        files={"file": ("f.xlsx", io.BytesIO(buf.getvalue()),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    sheet_id = response.json()[0]["id"]
    cells = (await client.get(f"/api/v1/sheets/{sheet_id}/cells")).json()
    by = {(c["row"], c["column"]): c for c in cells}
    assert by[(2, 0)]["formula"] == "=SUM(A1:A2)"
    assert by[(2, 0)]["value"] == "30"  # recalc fires


@pytest.mark.asyncio
async def test_xlsx_export_roundtrip(client):
    sheet_id = await _make_sheet(client)
    await client.patch(
        f"/api/v1/sheets/{sheet_id}/cells",
        json={
            "cells": [
                {"row": 0, "column": 0, "value": "name"},
                {"row": 0, "column": 1, "value": "qty"},
                {"row": 1, "column": 0, "value": "apple"},
                {"row": 1, "column": 1, "value": "3", "data_type": "number"},
            ]
        },
    )
    response = await client.get(f"/api/v1/sheets/{sheet_id}/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws["A1"].value == "name"
    assert ws["B1"].value == "qty"
    assert ws["A2"].value == "apple"
    assert ws["B2"].value == 3


@pytest.mark.asyncio
async def test_chart_recommendation_picks_bar_for_categorical(client):
    sheet_id = await _make_sheet(client)
    await client.patch(
        f"/api/v1/sheets/{sheet_id}/cells",
        json={
            "cells": [
                {"row": 0, "column": 0, "value": "label"},
                {"row": 0, "column": 1, "value": "value"},
                {"row": 1, "column": 0, "value": "a"},
                {"row": 1, "column": 1, "value": "10", "data_type": "number"},
                {"row": 2, "column": 0, "value": "b"},
                {"row": 2, "column": 1, "value": "20", "data_type": "number"},
                {"row": 3, "column": 0, "value": "c"},
                {"row": 3, "column": 1, "value": "30", "data_type": "number"},
            ]
        },
    )
    response = await client.get(
        f"/api/v1/sheets/{sheet_id}/chart",
        params={"start": "A1", "end": "B4", "has_header": True},
    )
    body = response.json()
    spec = body["vega_lite"]
    assert spec["mark"] == "bar"
    assert spec["encoding"]["x"]["field"] == "label"
    assert spec["encoding"]["y"]["field"] == "value"
    assert len(spec["data"]["values"]) == 3


@pytest.mark.asyncio
async def test_chart_recommendation_picks_line_for_dates(client):
    sheet_id = await _make_sheet(client)
    await client.patch(
        f"/api/v1/sheets/{sheet_id}/cells",
        json={
            "cells": [
                {"row": 0, "column": 0, "value": "date"},
                {"row": 0, "column": 1, "value": "revenue"},
                {"row": 1, "column": 0, "value": "2026-01-01"},
                {"row": 1, "column": 1, "value": "100", "data_type": "number"},
                {"row": 2, "column": 0, "value": "2026-02-01"},
                {"row": 2, "column": 1, "value": "150", "data_type": "number"},
                {"row": 3, "column": 0, "value": "2026-03-01"},
                {"row": 3, "column": 1, "value": "200", "data_type": "number"},
            ]
        },
    )
    response = await client.get(
        f"/api/v1/sheets/{sheet_id}/chart",
        params={"start": "A1", "end": "B4"},
    )
    spec = response.json()["vega_lite"]
    assert spec["mark"] == "line"
    assert spec["encoding"]["x"]["type"] == "temporal"


@pytest.mark.asyncio
async def test_chart_recommendation_picks_scatter_for_two_numerics(client):
    sheet_id = await _make_sheet(client)
    await client.patch(
        f"/api/v1/sheets/{sheet_id}/cells",
        json={
            "cells": [
                {"row": 0, "column": 0, "value": "x"},
                {"row": 0, "column": 1, "value": "y"},
                {"row": 1, "column": 0, "value": "1", "data_type": "number"},
                {"row": 1, "column": 1, "value": "2", "data_type": "number"},
                {"row": 2, "column": 0, "value": "3", "data_type": "number"},
                {"row": 2, "column": 1, "value": "5", "data_type": "number"},
                {"row": 3, "column": 0, "value": "5", "data_type": "number"},
                {"row": 3, "column": 1, "value": "8", "data_type": "number"},
            ]
        },
    )
    response = await client.get(
        f"/api/v1/sheets/{sheet_id}/chart",
        params={"start": "A1", "end": "B4"},
    )
    spec = response.json()["vega_lite"]
    assert spec["mark"] == "point"
