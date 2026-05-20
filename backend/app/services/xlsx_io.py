"""XLSX import / export via openpyxl. No formula evaluation here — values
flow through the formula engine on write, exports just emit raw values and
formulas.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from uuid import UUID

from openpyxl import Workbook as XlWorkbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Sheet
from app.repositories.cell_repo import CellRepository
from app.repositories.sheet_repo import SheetRepository
from app.schemas.cell import CellUpsert


def _infer_type(value) -> tuple[str, str]:
    """Return (data_type, serialised_value) for an openpyxl cell value."""
    if value is None:
        return "string", ""
    if isinstance(value, bool):
        return "boolean", "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return "number", str(int(value)) if float(value).is_integer() else str(value)
    if isinstance(value, (datetime, date)):
        return "date", value.isoformat()
    s = str(value)
    if s.startswith("="):
        return "formula", s
    return "string", s


async def import_xlsx(
    db: AsyncSession,
    workbook_id: UUID,
    sheet_name: str | None,
    xlsx_bytes: bytes,
) -> list[Sheet]:
    """Import every worksheet in the XLSX as its own Sheet. Returns created sheets."""
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    created: list[Sheet] = []

    for position, ws in enumerate(wb.worksheets):
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        row_count = max(len(rows), 100)
        column_count = max((len(r) for r in rows), default=0) or 26

        sheet = Sheet(
            workbook_id=workbook_id,
            name=sheet_name if (sheet_name and len(wb.worksheets) == 1) else ws.title,
            position=position,
            row_count=max(row_count, 100),
            column_count=max(column_count, 26),
        )
        sheet = await SheetRepository(db).create(sheet)
        created.append(sheet)

        cell_payloads: list[CellUpsert] = []
        for r_idx, row in enumerate(rows):
            for c_idx, raw in enumerate(row):
                data_type, value = _infer_type(raw)
                if data_type == "formula":
                    cell_payloads.append(
                        CellUpsert(row=r_idx, column=c_idx, value=None, formula=value, data_type="formula")
                    )
                else:
                    cell_payloads.append(
                        CellUpsert(row=r_idx, column=c_idx, value=value, data_type=data_type)
                    )

        if cell_payloads:
            await CellRepository(db).bulk_upsert(sheet.id, cell_payloads)

    return created


async def export_sheet_xlsx(db: AsyncSession, sheet_id: UUID) -> bytes:
    """Render a single sheet to a one-tab XLSX workbook."""
    sheet = await SheetRepository(db).get_by_id(sheet_id)
    if sheet is None:
        raise ValueError("Sheet not found")
    cells = await CellRepository(db).list_by_sheet(sheet_id)

    wb = XlWorkbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = sheet.name[:31]  # Excel sheet name limit

    for c in cells:
        # openpyxl uses 1-based row/column
        target = ws.cell(row=c.row + 1, column=c.column + 1)
        if c.formula:
            target.value = c.formula
        elif c.value is None or c.value == "":
            target.value = None
        else:
            target.value = _typed_export(c.value, c.data_type)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _typed_export(value: str, data_type: str | None) -> object:
    dt = (data_type or "").lower()
    if dt == "number":
        try:
            f = float(value)
            return int(f) if f.is_integer() else f
        except ValueError:
            return value
    if dt == "boolean":
        return value.upper() == "TRUE"
    return value
